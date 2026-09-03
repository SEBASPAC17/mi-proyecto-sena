from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


OUT = Path(__file__).resolve().parent / "AVANCE-PROYECTO-SENA-DINAMICASH.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(90, 90, 90)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13, color=BLUE if level == 1 else DARK_BLUE, bold=True)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.style = "Normal"
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        shade_cell(hdr[idx], header_fill)
        p = hdr[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, size=10, bold=True, color=DARK_BLUE)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            shade_cell(cells[idx], WHITE)
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            set_run_font(run, size=10)
    set_table_width(table, widths)
    return table


def add_metadata_table(doc):
    rows = [
        ("Proyecto", "Dinamicash Wallet"),
        ("Aprendiz", "Sebastian Andres Castaneda Panqueva"),
        ("Programa", "Analisis y Desarrollo de Software - SENA"),
        ("Tipo de entrega", "Revision de avance, no entrega final"),
        ("Fecha de preparacion", "31 de agosto de 2026"),
        ("Repositorio", "https://github.com/SEBASPAC17/mi-proyecto-sena.git"),
    ]
    table = add_table(doc, ["Campo", "Detalle"], rows, [2700, 6660], LIGHT_GRAY)
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    for style_name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True

    header = section.header.paragraphs[0]
    header.text = "Dinamicash Wallet - Informe de avance"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.runs[0], size=9, color=GRAY)

    footer = section.footer.paragraphs[0]
    footer.text = "Proyecto formativo SENA"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.runs[0], size=9, color=GRAY)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Informe de avance del proyecto")
    set_run_font(run, size=24, bold=True, color=RGBColor(0, 0, 0))

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Dinamicash Wallet - Revision tecnica de avance")
    set_run_font(run, size=13, color=GRAY)

    add_metadata_table(doc)

    add_heading(doc, "1. Resumen del proyecto")
    add_body(
        doc,
        "Dinamicash Wallet es una aplicacion web para la gestion de finanzas personales. "
        "Permite registrar ingresos, gastos, categorias, metas de ahorro y consultar reportes financieros."
    )
    add_body(
        doc,
        "En el avance actual implemente una arquitectura cliente-servidor con frontend en HTML, CSS y JavaScript, "
        "backend en Node.js con Express y persistencia de datos en MySQL."
    )

    add_heading(doc, "2. Objetivo del avance")
    add_body(
        doc,
        "Presentar los modulos funcionales que he desarrollado hasta la fecha y evidenciar que el aplicativo cuenta "
        "con versionamiento en GitHub, conexion a base de datos, autenticacion, API REST y pruebas basicas con Postman."
    )

    add_heading(doc, "3. Tecnologias utilizadas")
    tech_rows = [
        ("Frontend", "HTML5, CSS3, JavaScript, Chart.js, jsPDF"),
        ("Backend", "Node.js, Express"),
        ("Base de datos", "MySQL, base local llamada dinamicash"),
        ("Seguridad", "bcryptjs, cookies de sesion, validacion de correo, bloqueo por intentos"),
        ("Pruebas", "Postman para endpoints y pruebas automatizadas en Node.js"),
        ("Versionamiento", "Git y GitHub"),
    ]
    add_table(doc, ["Area", "Tecnologias / evidencia"], tech_rows, [2300, 7060])

    add_heading(doc, "4. Modulos desarrollados")
    modules = [
        "Registro de usuarios con validacion de datos.",
        "Inicio de sesion con PIN y sesion persistente.",
        "Verificacion de correo y recuperacion de PIN.",
        "Dashboard principal con resumen financiero.",
        "Gestion de movimientos: ingresos y gastos.",
        "Gestion de categorias y metas de ahorro.",
        "Reportes por periodo, tipo y categoria.",
        "Auditoria de eventos sensibles en la base de datos.",
        "Preparacion de APK Android mediante Capacitor.",
    ]
    for item in modules:
        add_bullet(doc, item)

    add_heading(doc, "5. Base de datos de prueba")
    add_body(
        doc,
        "La base de datos de prueba se llama dinamicash y esta implementada en MySQL. "
        "Contiene tablas para usuarios, categorias, movimientos, metas, tokens de correo, sesiones, tokens biometricos y auditoria."
    )
    db_rows = [
        ("Usuario demo", "demo.sena@dinamicash.com"),
        ("PIN demo", "1234"),
        ("Estado", "Usuario verificado y con autorizacion de datos"),
        ("Movimientos guardados", "6 registros de ingresos y gastos"),
        ("Metas guardadas", "2 metas de ahorro"),
        ("Tablas principales", "usuarios, categorias, movimientos, metas, user_sessions, audit_logs"),
    ]
    add_table(doc, ["Elemento", "Detalle"], db_rows, [2600, 6760], LIGHT_GRAY)

    doc.add_page_break()

    add_heading(doc, "6. Pruebas realizadas")
    add_body(
        doc,
        "Realice pruebas en Postman para validar que el backend responde correctamente y que los datos se consultan desde MySQL."
    )
    postman_rows = [
        ("POST /login", "Valida el ingreso del usuario demo", "200 OK"),
        ("GET /auth/session", "Comprueba la sesion activa", "200 OK"),
        ("GET /movimientos/1", "Lista movimientos guardados en la base de datos", "200 OK"),
        ("GET /metas/1", "Lista metas guardadas en la base de datos", "200 OK"),
        ("GET /reportes/1 (agosto 2026)", "Genera reporte con datos del periodo", "200 OK esperado"),
    ]
    add_table(doc, ["Endpoint", "Validacion", "Resultado"], postman_rows, [2600, 5060, 1700])

    add_heading(doc, "7. Evidencias disponibles")
    evidence = [
        ("GitHub", "Repositorio publico con commit visible: Avance proyecto SENA."),
        ("Aplicativo local", "Lo ejecute en http://localhost:3000/frontend/login.html."),
        ("Base de datos", "Consulta directa a MySQL con tablas y registros demo."),
        ("Postman", "Pruebas de login y consulta de movimientos con respuesta 200 OK."),
        ("Codigo fuente", "Carpetas frontend, backend, docs y configuracion Android."),
    ]
    add_table(doc, ["Evidencia", "Descripcion"], evidence, [2300, 7060], LIGHT_GRAY)

    add_heading(doc, "8. Lista de chequeo de avance")
    checklist_rows = [
        ("Versionamiento Git/GitHub", "Realizado", "Repositorio conectado y commit de avance visible."),
        ("Frontend funcional", "Realizado", "Pantallas de login, registro y dashboard disponibles."),
        ("Backend/API REST", "Realizado", "Endpoints de login, movimientos, metas y reportes."),
        ("Base de datos MySQL", "Realizado", "Base dinamicash con datos demo guardados."),
        ("Pruebas Postman", "Realizado", "Login y consulta de movimientos validados."),
        ("Datos de prueba", "Realizado", "Usuario demo con movimientos y metas."),
        ("Documento de avance", "Realizado", "Este informe consolida el estado del proyecto."),
        ("Entrega final", "Pendiente", "Ajustes finales, documentacion completa y validaciones ampliadas."),
    ]
    add_table(doc, ["Item", "Estado", "Observacion"], checklist_rows, [2800, 1700, 4860])

    add_heading(doc, "9. Pendientes y siguientes pasos")
    pending = [
        "Tomar capturas de evidencia para anexar o mostrar durante la revision.",
        "Organizar capturas y anexos de evidencia en caso de que sean solicitados durante la revision.",
        "Completar pruebas adicionales de endpoints si se solicita mayor evidencia tecnica.",
        "Pulir documentacion final y preparar sustentacion completa para la entrega definitiva.",
        "Corregir la configuracion del servicio MySQL96 de Windows, que actualmente apunta a un my.ini inexistente.",
    ]
    for item in pending:
        add_bullet(doc, item)

    add_heading(doc, "10. Guion corto para la revision")
    add_body(
        doc,
        "Dinamicash Wallet es una aplicacion de finanzas personales. En este avance presento que el sistema "
        "ya permite iniciar sesion, guardar movimientos y metas, consultar reportes y persistir datos en MySQL. "
        "Tambien se demuestra el versionamiento en GitHub y pruebas de API en Postman."
    )
    add_body(
        doc,
        "Durante la revision voy a mostrar primero el aplicativo, luego GitHub, despues Postman y por ultimo "
        "la consulta directa de la base de datos."
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
